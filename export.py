from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from io import BytesIO
from datetime import datetime
import re


def strip_inline_markdown(text):
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'\1', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'__(.+?)__', r'\1', text)
    text = re.sub(r'_(.+?)_', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    text = re.sub(r'~~(.+?)~~', r'\1', text)
    text = re.sub(r'^\s*>\s*', '', text, flags=re.MULTILINE)
    text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', text)
    text = text.strip()
    return text


def is_table_separator(line):
    stripped = line.strip()
    return bool(re.match(r'^[\|\s\-:]+$', stripped) and '|' in stripped)


def parse_markdown_table(lines):
    rows = []
    for line in lines:
        if is_table_separator(line):
            continue
        if '|' not in line:
            continue
        cells = [c.strip() for c in line.split('|')]
        cells = [c for c in cells if c != '']
        if cells:
            rows.append(cells)
    return rows


def add_table_to_doc(doc, rows):
    if not rows:
        return
    num_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = row.cells[j]
                cell.text = strip_inline_markdown(cell_text)
                if i == 0:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True
                            run.font.size = Pt(10)
                else:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(10)


def generate_docx(brief: str, query: str) -> bytes:
    doc = Document()

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # Title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run("Analyst Report")
    title_run.font.size = Pt(24)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(0x0F, 0x11, 0x17)

    # Query as subtitle
    clean_query = strip_inline_markdown(query)
    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_para.add_run(clean_query)
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Date
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_para.add_run(
        f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
    )
    date_run.font.size = Pt(10)
    date_run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    doc.add_paragraph()

    # Horizontal rule after header
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CCCCCC')
    pBdr.append(bottom)
    pPr.append(pBdr)

    lines = brief.split('\n')
    i = 0
    table_buffer = []
    in_table = False

    while i < len(lines):
        line = lines[i]
        raw = line.rstrip()

        # Detect markdown table start
        if '|' in raw and not is_table_separator(raw):
            next_line = lines[i + 1].rstrip() if i + 1 < len(lines) else ''
            if is_table_separator(next_line) or in_table:
                in_table = True
                table_buffer.append(raw)
                i += 1
                continue

        # Flush table when leaving table block
        if in_table and ('|' not in raw or raw.strip() == ''):
            if table_buffer:
                rows = parse_markdown_table(table_buffer)
                add_table_to_doc(doc, rows)
                doc.add_paragraph()
                table_buffer = []
            in_table = False

        if raw.strip() == '':
            i += 1
            continue

        # H1
        if re.match(r'^#\s+', raw):
            text = strip_inline_markdown(re.sub(r'^#+\s+', '', raw).strip())
            if text:
                heading = doc.add_heading(text, level=1)
                heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in heading.runs:
                    run.font.size = Pt(16)
                    run.bold = True
            i += 1
            continue

        # H2
        if re.match(r'^##\s+', raw):
            text = strip_inline_markdown(re.sub(r'^#+\s+', '', raw).strip())
            if text:
                doc.add_paragraph()
                heading = doc.add_heading(text, level=2)
                for run in heading.runs:
                    run.font.size = Pt(14)
                    run.bold = True
                    run.font.color.rgb = RGBColor(0x1D, 0x9E, 0x75)
            i += 1
            continue

        # H3
        if re.match(r'^###\s+', raw):
            text = strip_inline_markdown(re.sub(r'^#+\s+', '', raw).strip())
            if text:
                heading = doc.add_heading(text, level=3)
                for run in heading.runs:
                    run.font.size = Pt(12)
                    run.bold = True
            i += 1
            continue

        # Bullet point
        if re.match(r'^\s*[-*]\s+', raw):
            text = strip_inline_markdown(re.sub(r'^\s*[-*]\s+', '', raw).strip())
            if text:
                bullet = doc.add_paragraph(style='List Bullet')
                run = bullet.add_run(text)
                run.font.size = Pt(11)
            i += 1
            continue

        # Numbered list
        if re.match(r'^\s*\d+\.\s+', raw):
            text = strip_inline_markdown(re.sub(r'^\s*\d+\.\s+', '', raw).strip())
            if text:
                numbered = doc.add_paragraph(style='List Number')
                run = numbered.add_run(text)
                run.font.size = Pt(11)
            i += 1
            continue

        # Blockquote
        if re.match(r'^\s*>\s*', raw):
            text = strip_inline_markdown(re.sub(r'^\s*>\s*', '', raw).strip())
            if text:
                quote = doc.add_paragraph()
                quote.paragraph_format.left_indent = Inches(0.4)
                run = quote.add_run(text)
                run.font.size = Pt(11)
                run.italic = True
                run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            i += 1
            continue

        # Table separator — skip
        if is_table_separator(raw):
            i += 1
            continue

        # Orphan pipe line not in a table block
        if raw.strip().startswith('|') and not in_table:
            cells = [c.strip() for c in raw.split('|') if c.strip()]
            if cells:
                p = doc.add_paragraph()
                run = p.add_run(' | '.join(cells))
                run.font.size = Pt(11)
            i += 1
            continue

        # Regular paragraph — preserve inline bold
        para = doc.add_paragraph()
        parts = re.split(r'(\*\*[^*]+\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                clean = part[2:-2].strip()
                run = para.add_run(clean)
                run.bold = True
                run.font.size = Pt(11)
            else:
                clean = strip_inline_markdown(part)
                if clean:
                    run = para.add_run(clean)
                    run.font.size = Pt(11)

        i += 1

    # Flush any remaining table
    if table_buffer:
        rows = parse_markdown_table(table_buffer)
        add_table_to_doc(doc, rows)

    # Footer
    doc.add_paragraph()
    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run(
        "Generated by AnalystAgent — Powered by Claude | Not financial advice"
    )
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    footer_run.italic = True

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
